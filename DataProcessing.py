import docx
import os
import pandas as pd
import numpy as np

spam = os.listdir('E:\\同路人暑期学堂\\result')
# result=pd.DataFrame(columns=('id','name','grade','class','5','6','7','8','9','10','11','12','13'))
result = []
print(len(spam))
for i in spam:
    print(i)
    doc = docx.Document('E:\\同路人暑期学堂\\result\\{}'.format(i))
    tbl = doc.tables[0]
    result.append(tbl.cell(0, 2).text)
    result.append(tbl.cell(0, 4).text)
    result.append(tbl.cell(0, 6).text)
    result.append(tbl.cell(1, 2).text)
    result.append(tbl.cell(1, 4).text)
    result.append(tbl.cell(1, 6).text)
    result.append(tbl.cell(2, 2).text)
    result.append(tbl.cell(2, 6).text)
    result.append(tbl.cell(3, 2).text)
    result.append(tbl.cell(4, 2).text)
    result.append(tbl.cell(4, 6).text)
    result.append(tbl.cell(5, 2).text)
    result.append(tbl.cell(6, 1).text)

x = np.array(result)
x = x.reshape(-1, 13)
xx = pd.DataFrame(x)
xx.to_excel('E:\\同路人暑期学堂\\data0705.xls')
